"""Informe de cobros pendientes: lee un CSV de facturas y genera un Excel.

El Excel resultante tiene 4 hojas:
    - Resumen:        cifras clave (total pendiente, importe vencido, % vencido)
    - Antigüedad:     importe pendiente agrupado por tramos de retraso
    - Por cliente:    cuánto debe cada cliente, de mayor a menor
    - Detalle:        una fila por factura pendiente, de más a menos retraso

Uso:
    python src/informe_cobros.py
    python src/informe_cobros.py --fecha 2026-10-22
"""

import argparse
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

CARPETA_PROYECTO = Path(__file__).resolve().parent.parent

COLUMNAS_OBLIGATORIAS = {
    "num_factura",
    "cliente",
    "fecha_emision",
    "fecha_vencimiento",
    "total",
    "importe_cobrado",
}

# Tramos de antigüedad: (límite superior inclusive en días de retraso, nombre)
LIMITES_TRAMOS = [float("-inf"), 0, 30, 60, 90, float("inf")]
NOMBRES_TRAMOS = ["No vencida", "1-30 días", "31-60 días", "61-90 días", "Más de 90 días"]

# Columnas que se formatean como dinero o como fecha en el Excel
COLUMNAS_DINERO = {"total", "importe_cobrado", "importe_pendiente", "importe"}
COLUMNAS_FECHA = {"fecha_emision", "fecha_vencimiento"}


def cargar_facturas(ruta: Path) -> pd.DataFrame:
    """Lee el CSV y comprueba que tiene las columnas que necesitamos."""
    facturas = pd.read_csv(ruta, parse_dates=["fecha_emision", "fecha_vencimiento"])
    faltan = COLUMNAS_OBLIGATORIAS - set(facturas.columns)
    if faltan:
        raise ValueError(f"Al CSV le faltan estas columnas: {sorted(faltan)}")
    return facturas


def calcular_pendientes(facturas: pd.DataFrame, fecha_informe: pd.Timestamp) -> pd.DataFrame:
    """Devuelve solo las facturas con importe pendiente, con días de retraso y tramo."""
    facturas = facturas.copy()
    facturas["importe_pendiente"] = (facturas["total"] - facturas["importe_cobrado"]).round(2)

    # Pendiente = queda algo por cobrar (0.005 evita falsos positivos por decimales)
    pendientes = facturas[facturas["importe_pendiente"] > 0.005].copy()

    pendientes["dias_retraso"] = (fecha_informe - pendientes["fecha_vencimiento"]).dt.days
    pendientes["tramo"] = pd.cut(
        pendientes["dias_retraso"], bins=LIMITES_TRAMOS, labels=NOMBRES_TRAMOS
    )
    return pendientes.sort_values("dias_retraso", ascending=False)


def resumen_por_tramo(pendientes: pd.DataFrame) -> pd.DataFrame:
    """Importe pendiente y número de facturas en cada tramo de antigüedad."""
    return (
        pendientes.groupby("tramo", observed=False)
        .agg(facturas=("num_factura", "count"), importe=("importe_pendiente", "sum"))
        .round(2)
        .reset_index()
    )


def resumen_por_cliente(pendientes: pd.DataFrame) -> pd.DataFrame:
    """Cuánto debe cada cliente y cuál es su factura más atrasada."""
    return (
        pendientes.groupby("cliente")
        .agg(
            facturas=("num_factura", "count"),
            importe=("importe_pendiente", "sum"),
            max_dias_retraso=("dias_retraso", "max"),
        )
        .round(2)
        .sort_values("importe", ascending=False)
        .reset_index()
    )


def resumen_general(
    facturas: pd.DataFrame, pendientes: pd.DataFrame, fecha_informe: pd.Timestamp
) -> pd.DataFrame:
    """Cifras clave en formato concepto/valor."""
    total_pendiente = pendientes["importe_pendiente"].sum()
    vencido = pendientes.loc[pendientes["dias_retraso"] > 0, "importe_pendiente"].sum()
    porcentaje_vencido = vencido / total_pendiente if total_pendiente else 0

    filas = [
        ("Fecha del informe", fecha_informe.date()),
        ("Facturas en el fichero", len(facturas)),
        ("Facturas con importe pendiente", len(pendientes)),
        ("Importe total pendiente", round(total_pendiente, 2)),
        ("Importe ya vencido", round(vencido, 2)),
        ("% vencido sobre pendiente", round(porcentaje_vencido, 4)),
    ]
    return pd.DataFrame(filas, columns=["concepto", "valor"])


def dar_formato(hoja) -> None:
    """Cabecera en negrita, anchos de columna, formatos de dinero y fecha."""
    relleno = PatternFill("solid", fgColor="1F3864")
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno
        celda.alignment = Alignment(horizontal="center")
    hoja.freeze_panes = "A2"

    for columna in hoja.columns:
        nombre = str(columna[0].value)
        ancho = max(len(str(c.value)) if c.value is not None else 0 for c in columna)
        hoja.column_dimensions[columna[0].column_letter].width = min(max(ancho + 3, 12), 40)
        for celda in columna[1:]:
            if nombre in COLUMNAS_DINERO:
                celda.number_format = '#,##0.00 "€"'
            elif nombre in COLUMNAS_FECHA:
                celda.number_format = "DD/MM/YYYY"


def exportar_excel(hojas: dict[str, pd.DataFrame], ruta: Path) -> None:
    """Escribe cada DataFrame en su hoja y aplica el formato."""
    ruta.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(ruta, engine="openpyxl") as escritor:
        for nombre, tabla in hojas.items():
            tabla.to_excel(escritor, sheet_name=nombre, index=False)
            dar_formato(escritor.sheets[nombre])

        # El resumen mezcla textos, números y un porcentaje: se formatea a mano
        resumen = escritor.sheets["Resumen"]
        resumen.column_dimensions["A"].width = 34
        resumen.column_dimensions["B"].width = 18
        for fila in range(2, resumen.max_row + 1):
            concepto = resumen.cell(fila, 1).value
            celda_valor = resumen.cell(fila, 2)
            celda_valor.alignment = Alignment(horizontal="right")
            if concepto in ("Importe total pendiente", "Importe ya vencido"):
                celda_valor.number_format = '#,##0.00 "€"'
            elif concepto == "% vencido sobre pendiente":
                celda_valor.number_format = "0.0%"
            elif concepto == "Fecha del informe":
                celda_valor.number_format = "DD/MM/YYYY"


def main() -> None:
    parser = argparse.ArgumentParser(description="Genera el informe de cobros pendientes.")
    parser.add_argument(
        "--entrada", type=Path, default=CARPETA_PROYECTO / "datos" / "facturas_ejemplo.csv"
    )
    parser.add_argument(
        "--salida", type=Path, default=CARPETA_PROYECTO / "salida" / "informe_cobros.xlsx"
    )
    parser.add_argument(
        "--fecha",
        type=date.fromisoformat,
        default=date.today(),
        help="Fecha del informe en formato AAAA-MM-DD (por defecto, hoy).",
    )
    argumentos = parser.parse_args()
    fecha_informe = pd.Timestamp(argumentos.fecha)

    facturas = cargar_facturas(argumentos.entrada)
    pendientes = calcular_pendientes(facturas, fecha_informe)

    hojas = {
        "Resumen": resumen_general(facturas, pendientes, fecha_informe),
        "Antigüedad": resumen_por_tramo(pendientes),
        "Por cliente": resumen_por_cliente(pendientes),
        "Detalle": pendientes.drop(columns=["base_imponible", "iva", "fecha_cobro"], errors="ignore"),
    }
    exportar_excel(hojas, argumentos.salida)

    print(f"Facturas leídas:      {len(facturas)}")
    print(f"Facturas pendientes:  {len(pendientes)}")
    print(f"Total pendiente:      {pendientes['importe_pendiente'].sum():,.2f} €")
    print(f"Informe guardado en:  {argumentos.salida}")


if __name__ == "__main__":
    main()
